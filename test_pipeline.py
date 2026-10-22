"""Self-check for the lecture import / analysis pipeline. Run: .venv/bin/python test_pipeline.py
No network: LLM calls are stubbed."""
import os
import sys
import tempfile

os.environ["BLOCK1_DB"] = os.path.join(tempfile.mkdtemp(), "test.db")

import app.db as db

db.DB_PATH = os.environ["BLOCK1_DB"]
db.init_db()

from app.ingest import slides as SL


def _lecture(status_ocr="not_run", status_sum="not_started"):
    conn = db.get_conn()
    cur = conn.execute(
        "INSERT INTO lectures(title, filename, source, ocr_status, summary_status) VALUES('L','L.pdf','pdf',?,?)",
        (status_ocr, status_sum),
    )
    conn.commit()
    lid = cur.lastrowid
    conn.close()
    return lid


def _row(lid):
    conn = db.get_conn()
    r = conn.execute("SELECT * FROM lectures WHERE id=?", (lid,)).fetchone()
    conn.close()
    return r


def test_stuck_statuses_reset_on_startup():
    lid = _lecture("running", "generating")
    db.init_db()
    r = _row(lid)
    assert r["ocr_status"] == "error" and r["summary_status"] == "error", dict(r)


def test_malformed_summary_sets_error():
    from app.llm import summaries as S

    lid = _lecture()
    real = S.chat_json
    S.chat_json = lambda *a, **k: {"unexpected": 1}
    try:
        S.generate_summary(lid)
        assert False, "should have raised"
    except ValueError:
        pass
    finally:
        S.chat_json = real
    assert _row(lid)["summary_status"] == "error"


def test_pptx_groups_and_notes():
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    grp = s.shapes.add_group_shape()
    grp.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text_frame.text = "grouped heart text"
    s.notes_slide.notes_text_frame.text = "remember the apex"
    path = os.path.join(tempfile.mkdtemp(), "d.pptx")
    prs.save(path)
    text = SL.extract_pptx(path)[0][1]
    assert "grouped heart text" in text and "remember the apex" in text, text


def test_redundant_ocr_is_not_captioned():
    from app.llm import captions as C

    lid = _lecture()
    conn = db.get_conn()
    same = "cardiac output equals heart rate times stroke volume in the resting adult"
    conn.execute("INSERT INTO slides(lecture_id, slide_num, text, ocr_text) VALUES(?,1,?,?)", (lid, same, same))
    conn.execute(
        "INSERT INTO slides(lecture_id, slide_num, text, ocr_text) VALUES(?,2,'Frank Starling',"
        "'curve axes preload stroke volume ventricular filling')",
        (lid,),
    )
    conn.commit()
    conn.close()
    nums = [r["slide_num"] for r in C._slides_needing_captions(lid)]
    assert nums == [2], nums


def test_analysis_error_recorded_and_steps_isolated():
    import app.main as M
    import app.ingest.ocr as OCR
    import app.llm.captions as CAP

    lid = _lecture()
    calls = []

    def boom(_):
        raise RuntimeError("vision down")

    real = (OCR.run_ocr_for_lecture, CAP.generate_captions_for_lecture, M.generate_summary)
    OCR.run_ocr_for_lecture = boom
    CAP.generate_captions_for_lecture = lambda l, force=False: calls.append("captions")
    M.generate_summary = lambda l: calls.append("summary")
    try:
        M._analyze_lecture(lid)
    finally:
        OCR.run_ocr_for_lecture, CAP.generate_captions_for_lecture, M.generate_summary = real
    assert calls == ["captions", "summary"], "OCR failure must not skip later steps"
    assert "vision down" in _row(lid)["analysis_error"]


def test_delete_removes_files():
    from fastapi.testclient import TestClient
    import app.main as M
    from app.ingest.images import IMAGES_ROOT

    M.UPLOAD_DIR = tempfile.mkdtemp()
    lid = _lecture()
    src = os.path.join(M.UPLOAD_DIR, "L.pdf")
    open(src, "w").write("x")
    imgdir = os.path.join(IMAGES_ROOT, str(lid))
    os.makedirs(imgdir, exist_ok=True)
    open(os.path.join(imgdir, "page_1.jpg"), "w").write("x")
    assert TestClient(M.app).delete(f"/api/lectures/{lid}").status_code == 200
    assert not os.path.exists(src) and not os.path.exists(imgdir)


def test_upload_parses_once_and_force_imports_duplicates():
    import pymupdf
    from fastapi.testclient import TestClient
    import app.main as M

    M.UPLOAD_DIR = tempfile.mkdtemp()
    M._queue_analysis = lambda lid, force_captions=False: None  # no background work

    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Frank Starling mechanism relates preload to stroke volume in the ventricle")
    data = pdf.tobytes()

    parses = []
    real = SL.extract_pdf
    SL.extract_pdf = lambda p: (parses.append(p), real(p))[1]
    c = TestClient(M.app)
    try:
        r1 = c.post("/api/lectures/import", files={"files": ("a.pdf", data)}).json()
        assert len(r1["imported"]) == 1 and len(parses) == 1, (r1, parses)  # once, not twice
        r2 = c.post("/api/lectures/import", files={"files": ("b.pdf", data)}).json()
        assert r2["duplicates"] and r2["duplicates"][0]["title"], r2  # names the match
        r3 = c.post("/api/lectures/import", files={"files": ("b.pdf", data)}, data={"force": "true"}).json()
        assert len(r3["imported"]) == 1, r3
    finally:
        SL.extract_pdf = real


if __name__ == "__main__":
    tests = [v for k, v in sorted(globals().items()) if k.startswith("test_")]
    for t in tests:
        t()
        print(f"  ok  {t.__name__}")
    print(f"\n{len(tests)} passed")
    sys.exit(0)
